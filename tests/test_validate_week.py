import tempfile
import unittest
from pathlib import Path

from scripts.validate_week import validate_week


PROJECT = Path(__file__).resolve().parents[1]


class ValidateWeekTests(unittest.TestCase):
    def test_missing_data_stays_draft_and_is_not_zero_filled(self):
        with tempfile.TemporaryDirectory() as temp:
            week = Path(temp) / "2026-07-06_2026-07-12_周报数据"
            week.mkdir()
            result = validate_week(week, write_result=False)
            self.assertEqual(result["status"], "draft")
            self.assertIn("缺少目录：01_YouTube", result["errors"])
            self.assertNotIn(0, result["missing_files"])

    def test_template_has_registration_and_manual_meta_sheets(self):
        from openpyxl import load_workbook

        path = PROJECT / "data" / "模板" / "版权保护登记表.xlsx"
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, ["YouTube_CID", "Facebook_RM", "Meta_RM分析"])
            self.assertEqual(workbook["Meta_RM分析"]["E6"].value, "=IFERROR(C6/A6,0)")
        finally:
            workbook.close()

    def test_real_export_shapes_are_recognized_and_self_operated_rows_are_filtered(self):
        week = PROJECT / "data" / "周报数据" / "2026-07-06_2026-07-12_周报数据"
        result = validate_week(week, write_result=False)
        self.assertEqual(result["youtube"]["accounts_found"], 8)
        self.assertEqual(result["facebook"]["accounts_found"], 3)
        self.assertGreater(result["distribution"]["source_rows"], result["distribution"]["self_operated_rows"])
        self.assertEqual(result["distribution"]["filter_rule"], "机构包含“自营”")
        self.assertNotIn("账号汇总", "".join(result["missing_files"]))


if __name__ == "__main__":
    unittest.main()
