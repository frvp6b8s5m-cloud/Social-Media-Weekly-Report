from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parents[1]
DEFAULT_ROOT = PROJECT / "data" / "周报数据"
SCHEMA_PATH = PROJECT / "config" / "report_schema.json"
WEEK_RE = re.compile(r"^(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})_周报数据$")
VERSION_RE = re.compile(r"_v(\d+)$", re.IGNORECASE)


def load_schema() -> dict:
    return json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))


def parse_week_name(path: Path):
    match = WEEK_RE.match(path.name)
    if not match:
        return None
    return tuple(datetime.strptime(value, "%Y-%m-%d").date() for value in match.groups())


def discover_latest(root: Path) -> Path:
    candidates = []
    for path in root.iterdir() if root.exists() else []:
        parsed = parse_week_name(path) if path.is_dir() else None
        if parsed:
            candidates.append((parsed[0], path))
    if not candidates:
        raise FileNotFoundError(f"未找到标准周文件夹：{root}")
    return max(candidates, key=lambda item: item[0])[1]


def version_info(path: Path):
    match = VERSION_RE.search(path.stem)
    version = int(match.group(1)) if match else 1
    base = VERSION_RE.sub("", path.stem).lower()
    return base, path.suffix.lower(), version


def select_latest_versions(files: list[Path]):
    groups = defaultdict(list)
    for path in files:
        base, extension, version = version_info(path)
        groups[(str(path.parent).lower(), base, extension)].append((version, path))
    selected, history = [], []
    for values in groups.values():
        values.sort(key=lambda item: item[0])
        selected.append(values[-1][1])
        if len(values) > 1:
            history.append({
                "selected": values[-1][1].name,
                "versions": [path.name for _, path in values],
            })
    return selected, history


def tabular_files(folder: Path, supported: set[str], recursive: bool = False):
    if not folder.exists():
        return []
    paths = folder.rglob("*") if recursive else folder.iterdir()
    return [
        path for path in paths
        if path.is_file() and path.suffix.lower() in supported and not path.name.startswith("~$")
    ]


def read_csv_table(path: Path):
    raw = path.read_bytes()
    text = None
    used_encoding = None
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            text = raw.decode(encoding)
            used_encoding = encoding
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError(f"无法识别CSV编码：{path.name}")
    rows = list(csv.reader(text.splitlines()))
    header_index = 0
    if rows and rows[0] and rows[0][0].lower().startswith("sep="):
        header_index = 2
    headers = rows[header_index] if len(rows) > header_index else []
    data = [dict(zip(headers, row)) for row in rows[header_index + 1:] if any(value.strip() for value in row)]
    return headers, data, used_encoding


def clean_stem(path: Path):
    return VERSION_RE.sub("", path.stem)


def extract_dates(rows: list[dict]):
    found = []
    for row in rows:
        value = str(row.get("日期") or "")[:10]
        try:
            found.append(date.fromisoformat(value))
        except ValueError:
            continue
    return sorted(set(found))


def check_date_coverage(result: dict, label: str, rows: list[dict], start: date, end: date):
    dates = extract_dates(rows)
    if dates and (dates[0] != start or dates[-1] != end):
        result["warnings"].append(
            f"{label}实际数据日期为 {dates[0]} 至 {dates[-1]}，与周期 {start} 至 {end} 不一致"
        )


def validate_youtube(result: dict, platform_dir: Path, schema: dict, supported: set[str], start: date, end: date):
    summary = {"accounts_found": 0, "accounts_expected": len(schema["youtube_accounts"]), "files": {}}
    if not platform_dir.exists():
        result["errors"].append("缺少目录：01_YouTube")
        result["youtube"] = summary
        return
    actual_dirs = {path.name for path in platform_dir.iterdir() if path.is_dir()}
    for account in schema["youtube_accounts"]:
        if account not in actual_dirs:
            result["missing_accounts"].append(account)
            continue
        account_dir = platform_dir / account
        files, history = select_latest_versions(tabular_files(account_dir, supported, recursive=True))
        result["version_history"].extend(history)
        stems = {clean_stem(path): path for path in files}
        missing = [marker for marker in schema["youtube_export_files"] if marker not in stems]
        if missing:
            result["missing_files"].extend(f"{account}/{name}.csv" for name in missing)
            continue
        headers, table_rows, _ = read_csv_table(stems["表格数据"])
        missing_headers = [name for name in schema["youtube_table_headers"] if name not in headers]
        if missing_headers:
            result["errors"].append(f"{account}/表格数据 缺少字段：{', '.join(missing_headers)}")
        _, total_rows, _ = read_csv_table(stems["总计"])
        check_date_coverage(result, account, total_rows, start, end)
        summary["accounts_found"] += 1
        summary["files"][account] = {
            "content_rows": len(table_rows),
            "export_files": sorted(stems),
        }
    result["youtube"] = summary


def validate_facebook(result: dict, platform_dir: Path, schema: dict, supported: set[str], start: date, end: date):
    summary = {
        "accounts_found": 0,
        "accounts_expected": len(schema["facebook_accounts"]),
        "available_metrics": {},
        "optional_unavailable": {},
    }
    if not platform_dir.exists():
        result["errors"].append("缺少目录：02_Facebook")
        result["facebook"] = summary
        return
    actual_dirs = {path.name for path in platform_dir.iterdir() if path.is_dir()}
    for account in schema["facebook_accounts"]:
        if account not in actual_dirs:
            result["missing_accounts"].append(account)
            continue
        files, history = select_latest_versions(tabular_files(platform_dir / account, supported, recursive=True))
        result["version_history"].extend(history)
        stems = {clean_stem(path): path for path in files}
        missing = [marker for marker in schema["facebook_required_metrics"] if marker not in stems]
        if missing:
            result["missing_files"].extend(f"{account}/{name}.csv" for name in missing)
            continue
        for marker in schema["facebook_required_metrics"]:
            _, rows, _ = read_csv_table(stems[marker])
            check_date_coverage(result, f"{account}/{marker}", rows, start, end)
        summary["accounts_found"] += 1
        summary["available_metrics"][account] = sorted(stems)
        optional_missing = [marker for marker in schema["facebook_optional_metrics"] if marker not in stems]
        if optional_missing:
            summary["optional_unavailable"][account] = optional_missing
    result["facebook"] = summary


def number(value):
    try:
        return float(str(value or "0").replace(",", ""))
    except ValueError:
        return 0.0


def safe_rate(numerator: float, denominator: float):
    return numerator / denominator if denominator else None


def validate_distribution(result: dict, folder: Path, schema: dict, supported: set[str]):
    files, history = select_latest_versions(tabular_files(folder, supported, recursive=False))
    result["version_history"].extend(history)
    if not files:
        result["missing_files"].append("03_分销引流/用户数据.csv")
        return [], []
    combined, output_headers, valid_files = [], [], []
    for path in files:
        if path.suffix.lower() not in {".csv", ".tsv"}:
            continue
        headers, rows, _ = read_csv_table(path)
        if all(field in headers for field in schema["distribution_expected_metrics"]):
            output_headers = headers
            combined.extend(rows)
            valid_files.append(path.name)
    if not valid_files:
        result["errors"].append("分销用户数据缺少“机构”或漏斗指标字段")
        return [], []
    filter_field = schema["distribution_filter"]["field"]
    filter_text = schema["distribution_filter"]["contains"]
    own_rows = [row for row in combined if filter_text in str(row.get(filter_field) or "")]
    if not own_rows:
        result["errors"].append(f"分销用户数据中没有“{filter_field}包含{filter_text}”的记录")
        return output_headers, []
    metrics = ["推广资源数", "点击数", "激活用户数", "付费人数", "订单数", "付费金额($)", "付费金额-平台后($)", "分销收益金额($)"]
    totals = {metric: round(sum(number(row.get(metric)) for row in own_rows), 2) for metric in metrics}
    clicks = totals["点击数"]
    activations = totals["激活用户数"]
    payers = totals["付费人数"]
    orders = totals["订单数"]
    revenue = totals["付费金额($)"]
    result["distribution"] = {
        "source_files": valid_files,
        "source_rows": len(combined),
        "filter_rule": f"{filter_field}包含“{filter_text}”",
        "self_operated_rows": len(own_rows),
        "self_operated_institutions": len({row.get(filter_field) for row in own_rows}),
        "totals": totals,
        "recalculated": {
            "激活转化率": safe_rate(activations, clicks),
            "付费转化率": safe_rate(payers, activations),
            "客单价_按付费人数": safe_rate(revenue, payers),
            "平均订单金额": safe_rate(revenue, orders),
            "单次点击收入": safe_rate(revenue, clicks),
        },
        "drama_dimension_available": any(name in output_headers for name in ("剧目", "剧集名称")),
    }
    if clicks == 0 and activations > 0:
        result["warnings"].append("自营数据的点击数为0但激活用户大于0；激活率和单次点击收入标记为“无法计算”，不按0处理")
    if not result["distribution"]["drama_dimension_available"]:
        result["warnings"].append("当前分销导出没有剧目字段；本周只能做自营账号/机构维度漏斗，不生成剧目排名")
    return output_headers, own_rows


def validate_rights(result: dict, rights_dir: Path, schema: dict, supported: set[str]):
    files, history = select_latest_versions(tabular_files(rights_dir, supported))
    result["version_history"].extend(history)
    register = next((path for path in files if "版权保护登记表" in path.stem and path.suffix.lower() == ".xlsx"), None)
    if register is None:
        result["missing_files"].append("04_版权保护/版权保护登记表.xlsx")
        return
    workbook = None
    try:
        workbook = load_workbook(register, read_only=False, data_only=False)
        counts = {}
        for sheet_name, expected_headers in schema["rights_sheets"].items():
            if sheet_name not in workbook.sheetnames:
                result["errors"].append(f"版权登记表缺少工作表：{sheet_name}")
                continue
            sheet = workbook[sheet_name]
            headers = [sheet.cell(6, col).value for col in range(1, 4)]
            if headers != expected_headers:
                result["errors"].append(f"{sheet_name} 第6行字段不符合模板")
            counts[sheet_name] = sum(
                1 for row in range(7, sheet.max_row + 1)
                if sheet.cell(row, 1).value and (sheet.cell(row, 2).value or sheet.cell(row, 3).value)
            )
        result["rights_registration"] = counts
        result["rights_file"] = register.name
        meta_name = schema["meta_manual_sheet"]
        if meta_name not in workbook.sheetnames:
            result["missing_files"].append(f"04_版权保护/版权保护登记表/{meta_name}")
            return
        meta = workbook[meta_name]
        matches_10k, blocked_10k = meta["A6"].value, meta["C6"].value
        fb_share, ig_share = meta["B10"].value, meta["B11"].value
        regions = [meta.cell(row, 2).value for row in range(16, 21) if meta.cell(row, 2).value]
        titles = [meta.cell(row, 7).value for row in range(16, 21) if meta.cell(row, 7).value]
        required = {
            "Meta匹配视频数量（万）": matches_10k,
            "Meta屏蔽视频数量（万）": blocked_10k,
            "Facebook匹配占比": fb_share,
            "Instagram匹配占比": ig_share,
            "Meta匹配地区排行": regions,
            "Meta匹配剧目Top 5": titles if len(titles) == 5 else None,
        }
        result["missing_fields"].extend(name for name, value in required.items() if value in (None, "", []))
        result["meta_rights"] = {
            "input_unit": "万",
            "matches_10k": matches_10k,
            "blocked_10k": blocked_10k,
            "matches": number(matches_10k) * 10000 if matches_10k not in (None, "") else None,
            "blocked": number(blocked_10k) * 10000 if blocked_10k not in (None, "") else None,
            "block_rate": safe_rate(number(blocked_10k), number(matches_10k)) if matches_10k not in (None, "") else None,
            "facebook_share": fb_share,
            "instagram_share": ig_share,
            "regions_provided": len(regions),
            "top_titles_provided": len(titles),
            "estimated": True,
            "region_metric": "匹配占比",
        }
        if fb_share is not None and ig_share is not None and abs(number(fb_share) + number(ig_share) - 1) > 0.001:
            result["warnings"].append("Facebook与Instagram匹配占比合计不等于100%")
    except Exception as exc:
        result["errors"].append(f"无法读取版权登记表：{exc}")
    finally:
        if workbook is not None:
            workbook.close()


def validate_week(week: Path, write_result: bool = True):
    schema = load_schema()
    supported = set(schema["supported_tabular_extensions"])
    parsed = parse_week_name(week)
    result = {
        "week_folder": week.name,
        "status": "draft",
        "errors": [],
        "warnings": [],
        "missing_accounts": [],
        "missing_files": [],
        "missing_fields": [],
        "version_history": [],
        "rights_registration": {},
        "analysis_policy": "事实 + 待验证假设 + 验证动作",
    }
    if not parsed:
        result["errors"].append("文件夹名称不符合 YYYY-MM-DD_YYYY-MM-DD_周报数据")
        start = end = date.min
    else:
        start, end = parsed
        if start.weekday() != 0 or (end - start).days != 6:
            result["errors"].append("日期范围必须为周一至周日，共7天")
        result["period"] = {"start": start.isoformat(), "end": end.isoformat()}

    validate_youtube(result, week / "01_YouTube", schema, supported, start, end)
    validate_facebook(result, week / "02_Facebook", schema, supported, start, end)
    distribution_headers, own_rows = validate_distribution(result, week / "03_分销引流", schema, supported)
    validate_rights(result, week / "04_版权保护", schema, supported)

    notes_dir = week / "05_业务补充"
    notes = [] if not notes_dir.exists() else [path for path in notes_dir.iterdir() if path.suffix.lower() in {".md", ".docx"}]
    if not notes:
        result["warnings"].append("未提供本周补充说明；业务事实和下周行动待补充")
    result["business_notes"] = [path.name for path in notes]

    if not result["errors"] and not result["missing_accounts"] and not result["missing_files"] and not result["missing_fields"]:
        result["status"] = "ready"

    if write_result:
        output_dir = PROJECT / "data" / "处理结果" / week.name
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        if own_rows and distribution_headers:
            with (output_dir / "分销自营账号.csv").open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=distribution_headers, extrasaction="ignore")
                writer.writeheader()
                writer.writerows(own_rows)
    return result


def print_human(result: dict):
    ready = result["status"] == "ready"
    print("=" * 58)
    print("　ReelShort 周报数据检查")
    print("=" * 58)
    print(f"状态：{'数据齐全，可生成周报' if ready else '已识别数据，仍有项目待补充'}")
    print(f"周期：{result.get('period', {}).get('start', '-')} 至 {result.get('period', {}).get('end', '-')}")
    youtube = result.get("youtube", {})
    facebook = result.get("facebook", {})
    print(f"YouTube：{youtube.get('accounts_found', 0)}/{youtube.get('accounts_expected', 8)} 个频道已识别")
    print(f"Facebook：{facebook.get('accounts_found', 0)}/{facebook.get('accounts_expected', 3)} 个主页已识别")
    distribution = result.get("distribution", {})
    if distribution:
        print(f"分销引流：共{distribution['source_rows']}条，筛出机构包含“自营”{distribution['self_operated_rows']}条")
    print(f"版权登记：{result.get('rights_file', '未提供')}")
    if result.get("rights_registration"):
        print(f"  YouTube CID {result['rights_registration'].get('YouTube_CID', 0)}部；Facebook RM {result['rights_registration'].get('Facebook_RM', 0)}部")
    for title, items in (("错误", result["errors"]), ("缺少账号", result["missing_accounts"]), ("缺少文件", result["missing_files"]), ("待填字段", result["missing_fields"]), ("说明", result["warnings"])):
        if items:
            print(f"\n{title}：")
            for item in items:
                print(f"  - {item}")
    print("\n详细结果已保存到 data\\处理结果")


def main():
    parser = argparse.ArgumentParser(description="校验ReelShort最新一周社媒数据")
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument("--week-path", type=Path)
    parser.add_argument("--no-write", action="store_true")
    parser.add_argument("--json", action="store_true", help="在控制台输出JSON而不是简明结果")
    args = parser.parse_args()
    try:
        week = args.week_path.resolve() if args.week_path else discover_latest(args.root.resolve())
        result = validate_week(week, write_result=not args.no_write)
        print(json.dumps(result, ensure_ascii=False, indent=2)) if args.json else print_human(result)
        return 0 if result["status"] == "ready" else 2
    except Exception as exc:
        print(json.dumps({"status": "error", "message": str(exc)}, ensure_ascii=False, indent=2))
        return 1


if __name__ == "__main__":
    sys.exit(main())
