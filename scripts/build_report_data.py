from __future__ import annotations

import csv
import json
import math
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parents[1]
WEEK = PROJECT / "data" / "周报数据" / "2026-07-06_2026-07-12_周报数据"
OUTPUT = PROJECT / "data" / "处理结果" / WEEK.name / "report_data.json"
START = date(2026, 7, 6)
END = date(2026, 7, 12)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def number(value):
    try:
        return float(str(value or "0").replace(",", "").replace("%", ""))
    except (TypeError, ValueError):
        return 0.0


def safe_div(numerator, denominator):
    return numerator / denominator if denominator else None


def read_csv(path: Path, encoding="utf-8-sig", header_index=0):
    rows = list(csv.reader(path.read_text(encoding=encoding).splitlines()))
    headers = rows[header_index]
    return [dict(zip(headers, row)) for row in rows[header_index + 1:] if any(cell.strip() for cell in row)]


def parse_publish_date(value):
    text = str(value or "").strip()
    if not text:
        return None
    for candidate in (text, text.replace("Z", "+00:00")):
        try:
            return datetime.fromisoformat(candidate).date()
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%b %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def extract_youtube_channel_name(account_dir: Path) -> str:
    for child in account_dir.iterdir():
        if not child.is_dir():
            continue
        match = re.search(r"\d{4}-\d{2}-\d{2}_\d{4}-\d{2}-\d{2}\s+(.+)$", child.name)
        if match and match.group(1).strip():
            return match.group(1).strip()
    return account_dir.name.removeprefix("YT_")


FACEBOOK_PAGE_NAMES = {
    "FB_英语": "ReelShort TV",
    "FB_泰语": "ReelShort Thai",
    "FB_印尼语": "ReelShort Indonesia",
}


def build_youtube():
    accounts = []
    contents = []
    for account_dir in sorted((WEEK / "01_YouTube").iterdir(), key=lambda path: path.name):
        if not account_dir.is_dir():
            continue
        channel_name = extract_youtube_channel_name(account_dir)
        table_path = next(account_dir.rglob("表格数据.csv"))
        total_path = next(account_dir.rglob("总计.csv"))
        rows = read_csv(table_path)
        daily_rows = read_csv(total_path)
        total_row = next((row for row in rows if row.get("内容") == "总计"), None)
        content_rows = [row for row in rows if row.get("内容") != "总计"]
        active_ids = set()
        published_ids = set()
        account_contents = []
        for row in content_rows:
            content_id = row.get("内容", "")
            if content_id:
                active_ids.add(content_id)
            publish_date = parse_publish_date(row.get("视频发布时间"))
            if content_id and publish_date and START <= publish_date <= END:
                published_ids.add(content_id)
            views = number(row.get("观看次数"))
            revenue = number(row.get("估算收入 (USD)"))
            item = {
                "id": content_id,
                "account": channel_name,
                "title": row.get("视频标题") or "未命名内容",
                "published_at": row.get("视频发布时间"),
                "views": views,
                "watch_hours": number(row.get("观看时长（小时）")),
                "subscribers": number(row.get("订阅人数")),
                "revenue": revenue,
                "impressions": number(row.get("展示次数")),
                "ctr": number(row.get("展示点击率 (%)")) / 100,
                "rpm": safe_div(revenue * 1000, views),
            }
            contents.append(item)
            account_contents.append(item)
        views = number(total_row.get("观看次数")) if total_row else sum(item["views"] for item in account_contents)
        revenue = number(total_row.get("估算收入 (USD)")) if total_row else sum(item["revenue"] for item in account_contents)
        daily_views = sum(number(row.get("观看次数")) for row in daily_rows)
        accounts.append({
            "account": channel_name,
            "active_content": len(active_ids),
            "published": len(published_ids),
            "views": views,
            "watch_hours": number(total_row.get("观看时长（小时）")) if total_row else sum(item["watch_hours"] for item in account_contents),
            "subscribers": number(total_row.get("订阅人数")) if total_row else sum(item["subscribers"] for item in account_contents),
            "revenue": revenue,
            "impressions": number(total_row.get("展示次数")) if total_row else sum(item["impressions"] for item in account_contents),
            "rpm": safe_div(revenue * 1000, views),
            "revenue_per_active_content": safe_div(revenue, len(active_ids)),
            "daily_views_total": daily_views,
            "views_reconciled": abs(views - daily_views) < 0.5,
        })
    total_views = sum(item["views"] for item in accounts)
    total_revenue = sum(item["revenue"] for item in accounts)
    ranked_by_views = sorted([item for item in contents if item["views"] > 0], key=lambda item: item["views"])
    threshold = ranked_by_views[max(0, math.floor(len(ranked_by_views) * 0.75) - 1)]["views"] if ranked_by_views else 0
    overall_rpm = safe_div(total_revenue * 1000, total_views)
    anomaly_pool = [item for item in contents if item["views"] >= threshold and (item["rpm"] or 0) < (overall_rpm or 0)]
    anomalies = sorted(anomaly_pool, key=lambda item: (item["rpm"] or 0, -item["views"]))[:5]
    return {
        "accounts": sorted(accounts, key=lambda item: item["revenue"], reverse=True),
        "summary": {
            "channels": len(accounts),
            "published": sum(item["published"] for item in accounts),
            "active_content": sum(item["active_content"] for item in accounts),
            "views": total_views,
            "watch_hours": sum(item["watch_hours"] for item in accounts),
            "subscribers": sum(item["subscribers"] for item in accounts),
            "revenue": total_revenue,
            "impressions": sum(item["impressions"] for item in accounts),
            "rpm": overall_rpm,
            "all_accounts_reconciled": all(item["views_reconciled"] for item in accounts),
        },
        "top_content_by_revenue": sorted(contents, key=lambda item: item["revenue"], reverse=True)[:5],
        "top_content_by_views": sorted(contents, key=lambda item: item["views"], reverse=True)[:5],
        "high_view_low_rpm": anomalies,
    }


def build_facebook():
    accounts = []
    for account_dir in sorted((WEEK / "02_Facebook").iterdir(), key=lambda path: path.name):
        if not account_dir.is_dir():
            continue
        item = {"account": FACEBOOK_PAGE_NAMES.get(account_dir.name, account_dir.name), "revenue_available": False}
        for metric_path in account_dir.glob("*.csv"):
            rows = read_csv(metric_path, encoding="utf-16", header_index=2)
            if metric_path.stem == "收入":
                item["revenue_available"] = True
                item["revenue"] = sum(number(row.get("content_monetization")) for row in rows)
                item["total_creator_income"] = sum(number(row.get("Primary")) for row in rows)
                item["subscriptions"] = sum(number(row.get("subscriptions")) for row in rows)
                item["stars"] = sum(number(row.get("stars")) for row in rows)
            else:
                item[metric_path.stem] = sum(number(row.get("Primary")) for row in rows)
        views = item.get("浏览量")
        item["rpm"] = safe_div(item.get("revenue", 0) * 1000, views) if item["revenue_available"] else None
        accounts.append(item)
    available = [item for item in accounts if item["revenue_available"]]
    return {
        "accounts": sorted(accounts, key=lambda item: item.get("revenue", -1), reverse=True),
        "summary": {
            "pages": len(accounts),
            "views": sum(item.get("浏览量", 0) for item in accounts),
            "reach": sum(item.get("浏览人数", 0) for item in accounts),
            "engagement": sum(item.get("互动", 0) for item in accounts),
            "link_clicks": sum(item.get("链接点击量", 0) for item in accounts),
            "revenue": sum(item.get("revenue", 0) for item in available),
            "revenue_pages": len(available),
            "revenue_missing_pages": [item["account"] for item in accounts if not item["revenue_available"]],
        },
    }


def build_distribution():
    rows = read_csv(WEEK / "03_分销引流" / "用户数据.csv")
    own = [row for row in rows if "自营" in str(row.get("机构") or "")]
    fields = ["推广资源数", "点击数", "激活用户数", "付费人数", "订单数", "付费金额($)", "付费金额-平台后($)", "分销收益金额($)"]
    totals = {field: sum(number(row.get(field)) for row in own) for field in fields}
    account_rows = []
    for row in own:
        account_rows.append({
            "account": row.get("账号名称") or "未命名账号",
            "institution": row.get("机构") or "",
            "resources": number(row.get("推广资源数")),
            "clicks": number(row.get("点击数")),
            "activations": number(row.get("激活用户数")),
            "payers": number(row.get("付费人数")),
            "orders": number(row.get("订单数")),
            "revenue": number(row.get("付费金额($)")),
            "net_after_platform": number(row.get("付费金额-平台后($)")),
        })
    revenue = totals["付费金额($)"]
    clicks = totals["点击数"]
    activations = totals["激活用户数"]
    payers = totals["付费人数"]
    orders = totals["订单数"]
    return {
        "summary": {
            "source_rows": len(rows),
            "self_operated_rows": len(own),
            "self_operated_institutions": len({row.get("机构") for row in own}),
            "resources": totals["推广资源数"],
            "clicks": clicks,
            "activations": activations,
            "activation_rate": safe_div(activations, clicks),
            "payers": payers,
            "payment_rate": safe_div(payers, activations),
            "orders": orders,
            "revenue": revenue,
            "net_after_platform": totals["付费金额-平台后($)"],
            "distribution_commission": totals["分销收益金额($)"],
            "arppu": safe_div(revenue, payers),
            "average_order_value": safe_div(revenue, orders),
            "revenue_per_click": safe_div(revenue, clicks),
            "drama_dimension_available": False,
        },
        "top_accounts_by_revenue": sorted(account_rows, key=lambda item: item["revenue"], reverse=True)[:5],
        "accounts": account_rows,
    }


def build_rights():
    workbook = load_workbook(WEEK / "04_版权保护" / "版权保护登记表.xlsx", data_only=False, read_only=False)
    try:
        registrations = {}
        for sheet_name in ("YouTube_CID", "Facebook_RM"):
            sheet = workbook[sheet_name]
            registrations[sheet_name] = [
                {
                    "title": sheet.cell(row, 1).value,
                    "online_at": sheet.cell(row, 2).value.isoformat() if hasattr(sheet.cell(row, 2).value, "isoformat") else str(sheet.cell(row, 2).value),
                    "completed_at": sheet.cell(row, 3).value.isoformat() if hasattr(sheet.cell(row, 3).value, "isoformat") else str(sheet.cell(row, 3).value),
                }
                for row in range(7, sheet.max_row + 1)
                if sheet.cell(row, 1).value and sheet.cell(row, 2).value and sheet.cell(row, 3).value
            ]
        meta = workbook["Meta_RM分析"]
        regions = [
            {"rank": meta.cell(row, 1).value, "region": meta.cell(row, 2).value, "share": meta.cell(row, 3).value}
            for row in range(16, 21) if meta.cell(row, 2).value
        ]
        titles = [
            {"rank": meta.cell(row, 6).value, "title": meta.cell(row, 7).value, "matches": meta.cell(row, 8).value}
            for row in range(16, 21) if meta.cell(row, 7).value
        ]
        matches_10k = number(meta["A6"].value)
        blocked_10k = number(meta["C6"].value)
        return {
            "youtube": registrations["YouTube_CID"],
            "facebook": registrations["Facebook_RM"],
            "meta": {
                "matches_10k": matches_10k,
                "blocked_10k": blocked_10k,
                "matches": matches_10k * 10000,
                "blocked": blocked_10k * 10000,
                "block_rate": safe_div(blocked_10k, matches_10k),
                "facebook_share": meta["B10"].value,
                "instagram_share": meta["B11"].value,
                "regions": regions,
                "top_titles": titles,
                "estimated": True,
            },
        }
    finally:
        workbook.close()


def build_notes():
    text = (WEEK / "05_业务补充" / "本周补充说明.md").read_text(encoding="utf-8")
    return {"raw_markdown": text}


def main():
    youtube = build_youtube()
    facebook = build_facebook()
    distribution = build_distribution()
    rights = build_rights()
    ad_revenue = youtube["summary"]["revenue"] + facebook["summary"]["revenue"]
    total_revenue = ad_revenue + distribution["summary"]["revenue"]
    data = {
        "report": {
            "week_id": "2026-W28",
            "period_start": START.isoformat(),
            "period_end": END.isoformat(),
            "currency": "USD",
            "status": "review_draft",
            "is_first_snapshot": True,
            "updated_at": datetime.now().astimezone().isoformat(timespec="minutes"),
            "revenue": {
                "total_available_scope": total_revenue,
                "ad_available_scope": ad_revenue,
                "youtube": youtube["summary"]["revenue"],
                "facebook_available_scope": facebook["summary"]["revenue"],
                "referral_gross_payment": distribution["summary"]["revenue"],
                "ad_share": safe_div(ad_revenue, total_revenue),
                "referral_share": safe_div(distribution["summary"]["revenue"], total_revenue),
            },
            "data_boundaries": [
                "首次汇报，无上周快照及最近四周趋势，不计算环比。",
                "Facebook收入仅包含英语主页已提供的content_monetization；泰语、印尼语收入未提供，不按0处理。",
                "分销数据仅使用机构字段包含“自营”的记录；导出无剧目字段，不做剧目排名。",
                "自营数据点击数为0但存在激活，激活率和单次点击收入无法计算。",
            ],
        },
        "youtube": youtube,
        "facebook": facebook,
        "distribution": distribution,
        "rights": rights,
        "notes": build_notes(),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(data, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
